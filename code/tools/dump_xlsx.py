# -*- coding: utf-8 -*-
"""导出 Excel 全部 sheet 内容为文本（保留行结构）。
用法: python dump_xlsx.py <xlsx路径> <输出txt路径> [max_rows]
"""
import sys, os

def dump_xlsx(path, out, max_rows=None):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    lines = []
    for sn in wb.sheetnames:
        ws = wb[sn]
        lines.append(f"===== SHEET: {sn} =====")
        r = 0
        for row in ws.iter_rows(values_only=True):
            r += 1
            if max_rows and r > max_rows:
                lines.append(f"...(截断，共 {ws.max_row} 行)")
                break
            cells = ["" if v is None else str(v) for v in row]
            # 去掉尾部空列
            while cells and cells[-1] == "":
                cells.pop()
            lines.append(f"R{r}: " + " | ".join(cells))
        lines.append("")
    wb.close()
    with open(out, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    print(f"OK {os.path.basename(path)} -> {os.path.basename(out)} sheets={wb.sheetnames}")

if __name__ == "__main__":
    p, o = sys.argv[1], sys.argv[2]
    mr = int(sys.argv[3]) if len(sys.argv) > 3 else None
    dump_xlsx(p, o, mr)
