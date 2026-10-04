# -*- coding: utf-8 -*-
"""处理舟数据大文件：结构摘要 + 前 N 行 dump（xlsx/xlsm/xlsb）。"""
import sys, os

def dump_xlsx(path, out, max_rows=80):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    lines = []
    for sn in wb.sheetnames:
        ws = wb[sn]
        lines.append(f"===== SHEET: {sn} (rows~{ws.max_row} cols~{ws.max_column}) =====")
        r = 0
        for row in ws.iter_rows(values_only=True):
            r += 1
            if r > max_rows:
                lines.append(f"...(截断，共 {ws.max_row} 行)")
                break
            cells = ["" if v is None else str(v)[:60] for v in row]
            while cells and cells[-1] == "":
                cells.pop()
            lines.append(f"R{r}: " + " | ".join(cells))
        lines.append("")
    wb.close()
    with open(out, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    print(f"OK {os.path.basename(path)} -> {os.path.basename(out)}")


def dump_xlsb(path, out, max_rows=80):
    import pyxlsb
    with pyxlsb.open_workbook(path) as wb:
        lines = []
        for sn in wb.sheets:
            lines.append(f"===== SHEET: {sn} =====")
            r = 0
            with wb.get_sheet(sn) as ws:
                for row in ws.rows():
                    r += 1
                    if r > max_rows:
                        lines.append("...(截断)")
                        break
                    cells = ["" if v is None else str(v)[:60] for v in row]
                    while cells and cells[-1] == "":
                        cells.pop()
                    lines.append(f"R{r}: " + " | ".join(cells))
            lines.append("")
    with open(out, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    print(f"OK {os.path.basename(path)} -> {os.path.basename(out)}")


if __name__ == "__main__":
    src, out = sys.argv[1], sys.argv[2]
    mr = int(sys.argv[3]) if len(sys.argv) > 3 else 80
    if src.lower().endswith(".xlsb"):
        dump_xlsb(src, out, mr)
    else:
        dump_xlsx(src, out, mr)
