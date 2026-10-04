# -*- coding: utf-8 -*-
"""probe_jingchen7.py — 净尘七轴：导出"第一次接手前的机械操作段" + 探测候选参考录像

## 核对标准（用户 2026-09-15）
> **第一次接手前的机械操作**（即赛博塑料空行会停止操作、由玩家接手之前的部分）
> 应与视频开头对应步骤的操作完全相符。

⇒ 本脚本：① 读 `净尘1000有模组.xlsm`，找到**第一个空行**，导出其前的全部操作；
② 探测 `VideoRecords\\净罪\\` 下的候选录像（参数 + 抽帧），看哪份能与之对应。
"""
from __future__ import annotations

import json
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

XLSM = r'E:\__file\舟数据\关卡\净尘1000有模组.xlsm'
VID_DIR = r'E:\MUMU共享文件夹\VideoRecords\净罪'
TMP = r'C:\Users\23922\Documents\ds_harness\ak\.tmp'


def dump_mech():
    import openpyxl
    wb = openpyxl.load_workbook(XLSM, read_only=True, data_only=True)
    ws = wb['作战记录']
    rows = list(ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True))
    header = [str(c) if c is not None else '' for c in rows[0]]
    mech, stop_row = [], None
    for i, r in enumerate(rows[1:], start=2):
        v = list(r) + [None] * 14
        if all(x is None or str(x).strip() == '' for x in v[:8]):
            stop_row = i
            break
        mech.append({'row': i, '费用': v[1], '帧数': v[2], '绝对时间': v[3],
                     '操作': v[4], '干员': v[5], '坐标': v[6], '朝向': v[7]})
    wb.close()
    print(f'=== 净尘1000有模组：机械操作段（第一次接手前的连续机器操作）===')
    print(f'空行（接手点）位于第 {stop_row} 行 ⇒ 机械段 = 行 2~{mech[-1]["row"]}，共 {len(mech)} 条\n')
    print(f'{"行":>3} {"费用":>5} {"帧数":>5} {"间隔":>7} {"操作":<5} {"干员":<12} {"坐标":<5} {"朝向":<4}')
    for m in mech:
        print(f'{m["row"]:>3} {str(m["费用"] or ""):>5} {str(m["帧数"] or ""):>5} '
              f'{str(m["绝对时间"] or ""):>7} {str(m["操作"] or ""):<5} {str(m["干员"] or ""):<12} '
              f'{str(m["坐标"] or ""):<5} {str(m["朝向"] or ""):<4}')
    out = os.path.join(TMP, 'jingchen7_mech_segment.json')
    json.dump({'xlsm': XLSM, 'header': header, 'stop_row': stop_row, 'mech_ops': mech},
              open(out, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    print(f'\n已导出核对清单：{out}')
    return mech


def probe_videos():
    import cv2
    if not os.path.isdir(VID_DIR):
        print('视频目录不存在:', VID_DIR); return
    files = []
    for root, _d, fs in os.walk(VID_DIR):
        for f in fs:
            if f.lower().endswith(('.mp4', '.mkv', '.flv')):
                files.append(os.path.join(root, f))
    print(f'\n=== 候选录像（{len(files)} 个）===')
    for p in files[:6]:
        cap = cv2.VideoCapture(p)
        if not cap.isOpened():
            print(f'  打不开 {p}'); continue
        fps = cap.get(cv2.CAP_PROP_FPS)
        n = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
        w = int(cap.get(cv2.CAP_PROP_FRAME_WIDTH)); h = int(cap.get(cv2.CAP_PROP_FRAME_HEIGHT))
        dur = n / fps if fps else 0
        print(f'  {os.path.relpath(p, VID_DIR)}  {w}x{h} @{fps:.3f}fps  {n} 帧  {dur:.1f}s  '
              f'({os.path.getsize(p)/1e9:.2f} GB)')
        # 抽开头 3 帧
        for k, fr in enumerate((0, int(fps * 5), int(fps * 15))):
            cap.set(cv2.CAP_PROP_POS_FRAMES, fr)
            ok, img = cap.read()
            if ok:
                name = os.path.join(TMP, f'jc7_{os.path.basename(p)[:12]}_{k}_{fr}.png')
                cv2.imwrite(name, img)
        cap.release()
    print('  已抽开头帧到 .tmp（用于目视确认是否是该轴对应的录像）')


if __name__ == '__main__':
    dump_mech()
    probe_videos()
