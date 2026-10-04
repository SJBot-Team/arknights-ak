# -*- coding: utf-8 -*-
"""diag_sheet1.py — 打印作者 Sheet1 的原始行（含时间列），核对「龙落地」等里程碑。

用途：确认我们解析出的 ops 时间（29.0s 落龙）与作者自己 Sheet1 记录的时间（66.1s）哪个是"落地"。
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
sys.stdout.reconfigure(encoding='utf-8')

from openpyxl import load_workbook                    # noqa: E402

from ak_sim import scenario_h124 as S                 # noqa: E402


def main():
    wb = load_workbook(S.XLSM, data_only=True)
    print('工作表:', wb.sheetnames)
    for name in wb.sheetnames:
        ws = wb[name]
        print(f'\n===== {name}  {ws.max_row} 行 × {ws.max_column} 列 =====')
        if 'Sheet1' not in name and len(wb.sheetnames) > 1 and name != wb.sheetnames[0]:
            continue
        for r in range(1, min(ws.max_row, 60) + 1):
            vals = []
            for c in range(1, min(ws.max_column, 14) + 1):
                v = ws.cell(row=r, column=c).value
                if v is None:
                    vals.append('')
                elif isinstance(v, float):
                    vals.append(f'{v:.3f}'.rstrip('0').rstrip('.'))
                else:
                    vals.append(str(v).replace('\n', '⏎')[:26])
            if any(vals):
                print(f'  r{r:>3} | ' + ' | '.join(vals))


if __name__ == '__main__':
    main()
