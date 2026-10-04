# -*- coding: utf-8 -*-
"""
axis.py — 轴表（操作轴）格式与解析：模拟器与打图执行器共用

设计依据：
  - 赛博塑料 PRTS+ 轴表：以「费用读数 + 费用内帧偏移」为唯一时间参照
    （见 E:\\__file\\舟数据\\关卡\\h12-4\\赛博塑料h12-4.xlsm『使用须知』第三条）
  - PRTS+ 限制：费用条满/剿灭/行商等费用不单调时失效 → 本模块扩展绝对帧参照（墙钟/视频锚点）

轴表 JSON 结构：
{
  "level": "obt/hard/level_hard_12-04",
  "timebase": "cost",              # cost | frame
  "seed": 0,
  "ops": [
     {"cost": 11, "foff": 15, "op": "deploy", "unit": "令", "coord": "C12", "dir": "上",
      "module": "uniequip_002_ling", "skill_index": 2, "note": ""},
     {"cost": 20, "foff": 0,  "op": "deploy", "unit": "蕾缪安", "coord": "G9", "dir": "下"},
     {"foff": 4,  "op": "skill", "unit": "蕾缪安"},
     {"foff": 8,  "op": "retreat", "unit": "蕾缪安"},
     {"frame": 1234, "op": "deploy", ...}   # timebase=frame 时使用绝对帧
  ]
}
"""
from __future__ import annotations

import json
import os
import string
from dataclasses import dataclass, field, asdict

FPS = 30


# ---------------- 坐标 ----------------
# 变换 = (行轴来源, 行是否反向, 列轴来源, 列是否反向)
# 关卡 mapData: map[r][c]，r 对应客户端场景 x、c 对应场景 y（已与 MAA Tile-Pos 交叉验证）
COORD_TRANSFORMS = {
    'rc_normal':   dict(row_from='letter', row_flip=False, col_from='number', col_flip=False),
    'rc_rowflip':  dict(row_from='letter', row_flip=True,  col_from='number', col_flip=False),
    'rc_colflip':  dict(row_from='letter', row_flip=False, col_from='number', col_flip=True),
    'rc_bothflip': dict(row_from='letter', row_flip=True,  col_from='number', col_flip=True),
    'cr_normal':   dict(row_from='number', row_flip=False, col_from='letter', col_flip=False),
    'cr_rowflip':  dict(row_from='number', row_flip=True,  col_from='letter', col_flip=False),
    'cr_colflip':  dict(row_from='number', row_flip=False, col_from='letter', col_flip=True),
    'cr_bothflip': dict(row_from='number', row_flip=True,  col_from='letter', col_flip=True),
}


def coord_to_rc(coord: str, rows: int, cols: int, transform: str = 'rc_normal') -> tuple[int, int]:
    """'C12' → (row, col) 0-based。字母 A-Z 表示一个轴，数字 1-N 表示另一轴。"""
    coord = coord.strip().upper()
    letters = ''.join(ch for ch in coord if ch.isalpha())
    digits = ''.join(ch for ch in coord if ch.isdigit())
    if not letters or not digits:
        raise ValueError(f'bad coord: {coord}')
    letter_idx = 0
    for ch in letters:
        letter_idx = letter_idx * 26 + (ord(ch) - ord('A') + 1)
    letter_idx -= 1                     # 0-based
    num_idx = int(digits) - 1           # 0-based
    t = COORD_TRANSFORMS.get(transform, COORD_TRANSFORMS['rc_normal'])
    r = letter_idx if t['row_from'] == 'letter' else num_idx
    c = num_idx if t['col_from'] == 'number' else letter_idx
    if t['row_flip']:
        r = (rows - 1) - r
    if t['col_flip']:
        c = (cols - 1) - c
    return r, c


def calibrate_coord_transform(bmap, coords: list[str], melee_flags: list[bool] | None = None) -> tuple[str, float]:
    """用「部署合法性」自动标定坐标变换（返回最佳变换与命中率）。"""
    rows, cols = bmap.height, bmap.width
    best, best_score = 'rc_normal', -1.0
    for name in COORD_TRANSFORMS:
        ok = 0
        tot = 0
        for i, cd in enumerate(coords):
            try:
                r, c = coord_to_rc(cd, rows, cols, name)
            except Exception:
                continue
            tot += 1
            melee = melee_flags[i] if melee_flags else True
            if bmap.can_deploy(r, c, melee=melee):
                ok += 1
        score = ok / tot if tot else 0.0
        if score > best_score:
            best, best_score = name, score
    return best, best_score


# ---------------- 轴表 ----------------
@dataclass
class AxisOp:
    op: str                       # deploy / retreat / skill / pause / speed
    unit: str = ''                # 干员名或简称
    coord: str = ''               # 部署坐标（如 C12）
    dir: str = ''                 # 朝向
    cost: int | None = None       # 费用读数（timebase=cost）
    foff: int = 0                 # 费用内帧偏移（0..29）
    frame: int | None = None      # 绝对帧（timebase=frame）
    abs_seconds: float | None = None   # 表内「绝对时间」列换算的秒数（D + C/30），仅作参考/备选时基
    module: str | None = None
    module_level: int = 3
    skill_index: int = 0
    skill_level: int = 10
    elite: int = 2
    level: int = 90
    potential: int = 6
    trust: int = 200
    note: str = ''
    short: str = ''               # 简称（轴表里可能用简称指代单位）
    raw: dict = field(default_factory=dict)

    def to_dict(self):
        d = asdict(self)
        d.pop('raw', None)
        keep_abs = d.get('abs_seconds')
        d = {k: v for k, v in d.items() if v not in (None, '', 0, 3, 10, 2, 90, 6, 200) or k in ('op',)}
        if keep_abs is not None:
            d['abs_seconds'] = keep_abs          # 该字段即使是 3/10 等"常见默认值"也必须保留
        return d


@dataclass
class Axis:
    level: str
    ops: list[AxisOp] = field(default_factory=list)
    timebase: str = 'cost'
    seed: int = 0
    meta: dict = field(default_factory=dict)

    # ---- 存取 ----
    def save(self, path: str):
        os.makedirs(os.path.dirname(path) or '.', exist_ok=True)
        data = {
            'level': self.level, 'timebase': self.timebase, 'seed': self.seed,
            'meta': self.meta,
            'ops': [o.to_dict() for o in self.ops],
        }
        with open(path, 'w', encoding='utf-8') as f:
            json.dump(data, f, ensure_ascii=False, indent=1)
        return path

    @classmethod
    def load(cls, path: str) -> 'Axis':
        d = json.load(open(path, encoding='utf-8'))
        a = cls(level=d['level'], timebase=d.get('timebase', 'cost'), seed=d.get('seed', 0),
                meta=d.get('meta', {}))
        for o in d.get('ops', []):
            a.ops.append(AxisOp(**{k: v for k, v in o.items() if k in AxisOp.__dataclass_fields__}))
        return a

    def to_markdown(self) -> str:
        """人类可读呈现（便于手动修改/对照）。"""
        lines = [f'# 轴表 {self.level}（timebase={self.timebase}, seed={self.seed}）', '',
                 '| # | 费用 | 费用内帧 | 操作 | 单位 | 坐标 | 朝向 | 备注 |',
                 '|---|---|---|---|---|---|---|---|']
        for i, o in enumerate(self.ops):
            lines.append(f'| {i+1} | {o.cost if o.cost is not None else ""} | {o.foff} | {o.op} | '
                         f'{o.unit} | {o.coord} | {o.dir} | {o.note} |')
        return '\n'.join(lines)

    def to_csv(self) -> str:
        rows = ['idx,cost,foff,op,unit,coord,dir,note']
        for i, o in enumerate(self.ops):
            rows.append(f'{i+1},{o.cost if o.cost is not None else ""},{o.foff},{o.op},{o.unit},'
                        f'{o.coord},{o.dir},{o.note}')
        return '\n'.join(rows)

    # ---- 解析为绝对帧（模拟器/打图端共用换算） ----
    def resolve_frames(self, initial_cost: float = 10.0, cost_max: int = 99,
                       regen_per_sec: float = 1.0, unit_costs: dict | None = None,
                       spent_of: dict | None = None, refund_ratio: float = 0.5,
                       timebase: str = 'cost-relative') -> list[dict]:
        """把 (费用读数, 费用内帧) 解析为绝对帧。

        语义（与 PRTS+ 一致，2026-09-02 修正）：
            显示费用 = 初始费用 + 自然回复量 − 累计已消耗（撤退按 50% 返还）
            → 自然回复量 = 显示费用 − 初始费用 + **净累计消耗**
            → 绝对帧 = 自然回复量 × 每费帧数 + 费用内帧

        旧实现有两处问题：①漏掉「累计已消耗」，部署后的锚点整体偏早；
        ②把空费用列顺延上一读数作为召唤物行的锚点。两者均已修正。
        与 `code/vision/plugins/clock/cost_ruler.py::build_frame_schedule` 为同一公式，
        一致性由 `ak/code/vision/apps/test_ruler_clock.py::test_schedule` 交叉校验。
        """
        per_cost = FPS / max(1e-9, regen_per_sec)
        unit_costs = unit_costs or {}
        spent_of = spent_of or {}

        def _c(u: str) -> float:
            return float(spent_of.get(u, unit_costs.get(u, 0)) or 0)

        out: list[dict] = []
        spent = 0.0
        abs_frame = 0.0
        for o in self.ops:
            raw = abs_frame
            if timebase == 'abs-seconds' and o.abs_seconds is not None:
                # 备选时基：直接采信表内「绝对时间」列 D + C/30（费用条满后不可靠，仅供对照）
                raw = o.abs_seconds * FPS
            elif timebase == 'hybrid' and o.cost is not None:
                # 混合读法：有费用读数的行用「读数 + 记账」定位，其余行用「绝对时间」列
                raw = (float(o.cost) - initial_cost + spent) * per_cost + o.foff
            elif timebase == 'hybrid' and o.abs_seconds is not None:
                raw = o.abs_seconds * FPS
            elif o.frame is not None:
                raw = float(o.frame)
            elif o.cost is not None:
                raw = (float(o.cost) - initial_cost + spent) * per_cost + o.foff
            else:
                raw = abs_frame + o.foff
            clamped = None
            if raw < abs_frame - 1e-9:      # 顺序语义：锚点早于上一条 → 夹紧并记录
                clamped = raw
                raw = abs_frame
            abs_frame = raw
            out.append({'frame': int(round(abs_frame)), 'op': o.op, 'unit': o.unit,
                        'coord': o.coord, 'dir': o.dir, 'module': o.module,
                        'module_level': o.module_level, 'skill_index': o.skill_index,
                        'skill_level': o.skill_level, 'elite': o.elite, 'level': o.level,
                        'potential': o.potential, 'trust': o.trust, 'note': o.note,
                        'cost': o.cost, 'foff': o.foff, 'spent_before': round(spent, 3),
                        'abs_seconds': o.abs_seconds, 'timebase': timebase,
                        'clamped_from': (round(clamped, 3) if clamped is not None else None),
                        'raw': o.raw})
            if o.op == 'deploy':
                spent += _c(o.unit)
            elif o.op == 'retreat':
                spent -= _c(o.unit) * refund_ratio
        return out


# ---------------- 赛博塑料（PRTS+ xlsm）解析 ----------------
def parse_cyberplastic(path: str, sheet: str = '作战记录') -> Axis:
    """解析赛博塑料/PRTS+ 的 xlsm 作战记录表。

    列：A 当前执行 | B 费用 | C 帧数 | D 绝对时间 | E 操作 | F 干员 | G 坐标 | H 朝向 | I 简称 | J 备注
    
    时间语义（已判定）：D = 绝对游戏秒（截断显示），C = D 的小数部分帧（(D*30) mod 30）；
    费用单调自然回复 1/s，读数 = 累计生成 − 已消耗（部署/撤退影响）。
    """
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_vba=False)
    ws = wb[sheet]
    ops: list[AxisOp] = []
    level_code = ''
    for r in range(1, 20):
        k = ws.cell(row=r, column=13).value
        v = ws.cell(row=r, column=14).value
        if k == '关卡代号' and v:
            level_code = str(v)
    cur_cost = None
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        vals = list(row) + [None] * 12
        cur_exe, cost_v, foff_v, abs_v, op_v, unit_v, coord_v, dir_v, short_v, note_v = vals[:10]
        if not op_v:
            continue
        opname = str(op_v).strip()
        kind = {'部署': 'deploy', '技能': 'skill', '撤退': 'retreat', '暂停': 'pause'}.get(opname, opname)
        # 2026-09-02 修正：费用列 B 只在**显式给出读数**的行作为锚点；
        # 空 B 表示「按帧数 C 相对上一行计时」，不能把上一个读数顺延（否则召唤物行的锚点全错）。
        has_cost = cost_v not in (None, '')
        if has_cost:
            try:
                cur_cost = int(float(cost_v))
            except Exception:
                pass
        foff = int(float(foff_v)) if foff_v not in (None, '') else 0
        abs_s = float(abs_v) if abs_v not in (None, '') else None
        ops.append(AxisOp(
            op=kind, unit=str(unit_v or '').strip(),
            coord=str(coord_v or '').strip(), dir=str(dir_v or '').strip(),
            cost=cur_cost if has_cost else None, foff=foff,
            abs_seconds=(abs_s + foff / FPS) if abs_s is not None else None,
            short=str(short_v or '').strip(), note=str(note_v or '').strip(),
            raw={'abs_time': abs_s,
                 'abs_seconds': (abs_s + foff / FPS) if abs_s is not None else None,
                 'cost_recorded': has_cost,
                 'abs_frames': (int(round(abs_s * FPS + foff)) if abs_s is not None else None)},
        ))
    return Axis(level=level_code, ops=ops, timebase='cost',
                meta={'source': os.path.basename(path), 'solver': '赛博塑料/PRTS+ 轴表',
                      'note': '费用列 B 仅在显式读数行作为锚点；空 B 行按 C 相对上一行计时。'
                              'D 列绝对时间在费用条满后不可靠（见使用须知第三条），仅作提示。'})


if __name__ == '__main__':
    import sys
    sys.stdout.reconfigure(encoding='utf-8')
    ax = parse_cyberplastic(r'E:\__file\舟数据\关卡\h12-4\赛博塑料h12-4.xlsm')
    print(f'解析 {len(ax.ops)} 条操作，关卡代号={ax.level}')
    for o in ax.ops[:8]:
        print(f'  cost={o.cost} foff={o.foff} {o.op:<8} {o.unit:<8} {o.coord:<5} {o.dir}')
    print()
    resolved = ax.resolve_frames(initial_cost=10)
    print('解析为绝对帧（前 8 条）:')
    for r in resolved[:8]:
        print('  ', r['frame'], r['op'], r['unit'], r['coord'])
