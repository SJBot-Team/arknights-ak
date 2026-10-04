# -*- coding: utf-8 -*-
"""analyze_plan_feasibility.py — 逐条对照"PRTS+ 计划的时刻"与"费用何时才真的够"

## 目的
轴表的「绝对时间 + 帧数/30」是 **PRTS+ 计时器**给出**计划时刻**；
而游戏能否在该时刻真的部署，取决于**当时费用是否足够**。
本脚本用费用账本模型（初始 10、上限 20、+1/s、按部署逐个扣费）逐条算出：
  · 计划时刻的**当时费用**；
  · 该操作**需要**的费用；
  · 是否可行；不可行则给出**最早可行时刻**与**顺延量**。
⇒ 输出"它的计时口径 vs 游戏实际"的偏差表（这正是用户要的"观察费用满时的现象"的离线版）。

## 费用表
用**满配**口径（含令的模组：弦惊 18 而非 23）。
"""
from __future__ import annotations

import sys

sys.stdout.reconfigure(encoding='utf-8')

XLSM = r'E:\__file\舟数据\关卡\h12-4\赛博塑料h12-4.xlsm'
FPS = 30
INIT, CAP, RATE = 10.0, 20.0, 1.0
COST = {'令': 10, '蕾缪安': 20, '弦惊': 18, '麒麟R夜刀': 8, 'U-Official': 3,
        '傀影': 8, '艾拉': 10, '雷鸣地雷': 5}
# 会在战斗中被摧毁/进点而提升上限的事件（来自 §7.120 的模拟器实测）
CAP_EVENTS = [(48.97, 22.0), (53.03, 24.0)]


def cap_at(t):
    c = CAP
    for te, cv in CAP_EVENTS:
        if t >= te:
            c = cv
    return c


def main():
    import openpyxl
    wb = openpyxl.load_workbook(XLSM, read_only=True, data_only=True)
    ws = wb['作战记录']
    rows = list(ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True))
    ops = []
    for r in rows:
        vals = list(r) + [None] * 12
        _cur, cost_col, fr, t, op, unit = vals[:6]
        if op is None or t is None or fr is None:
            continue
        ops.append({'op': str(op), 'unit': str(unit or ''), 't': float(t),
                    'fr': int(fr), 'plan': float(t) + float(fr) / FPS,
                    'cost_col': cost_col})
    wb.close()
    ops.sort(key=lambda o: o['plan'])
    print(f'轴表操作 {len(ops)} 条（按计划时刻排序）\n')

    cur, last = INIT, 0.0
    print(f'{"#":>3} {"计划秒":>7} {"操作":<6} {"单位":<10} {"需要":>4} {"当时费用":>7} '
          f'{"可行":<4} {"最早可行":>7} {"顺延":>6}')
    total_delay = 0.0
    infeasible = 0
    for i, o in enumerate(ops, 1):
        # 累积费用到计划时刻
        cap = cap_at(o['plan'])
        cur = min(cap, cur + RATE * (o['plan'] - last))
        need = COST.get(o['unit'])
        if o['op'] in ('部署', 'deploy') and need:
            ok = cur >= need
            earliest = o['plan']
            if not ok:
                earliest = o['plan'] + (need - cur) / RATE
                infeasible += 1
                total_delay += earliest - o['plan']
            print(f'{i:>3} {o["plan"]:>7.2f} {o["op"]:<6} {o["unit"]:<10} {need:>4} '
                  f'{cur:>7.2f} {"是" if ok else "否":<4} {earliest:>7.2f} '
                  f'{(0 if ok else earliest-o["plan"]):>6.2f}')
            if ok:
                cur -= need
                last = o['plan']
            else:
                cur = min(cap_at(earliest), cur + RATE * (earliest - o['plan'])) - need
                last = earliest
        else:
            print(f'{i:>3} {o["plan"]:>7.2f} {o["op"]:<6} {o["unit"]:<10} {"-":>4} '
                  f'{cur:>7.2f} {"-":<4} {"-":>7} {"-":>6}')
            last = o['plan']
    print(f'\n不可行操作 {infeasible}/{len(ops)} 条，累计顺延 {total_delay:.2f} 秒')
    print('说明：PRTS+ 的"计划时刻"来自它自己的计时器（费用条每满一次记 1 秒、封顶期间停表）；')
    print('      若该时刻费用不足，游戏端会等待到费用够才真正落地 ⇒ 表格里的"顺延"就是差异。')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
