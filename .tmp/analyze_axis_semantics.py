# -*- coding: utf-8 -*-
"""临时分析：h12-4 赛博塑料轴表 时间/费用语义核验

模型 A：T = D + C/30（秒）；显示费用 = 初始费用 + T×回费速率 − 累计已消耗（撤退退 50%）
用途：用「有费用读数的行」检验模型，并找出偏差起点与偏差量。
"""
import sys
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl

P = r'E:\__file\舟数据\关卡\h12-4\赛博塑料h12-4.xlsm'
FPS = 30
INIT = 10
COST = {'令': 10, '蕾缪安': 20, '弦惊': 23, '麒麟R夜刀': 8, 'U-Official': 3,
        '傀影': 8, '艾拉': 10, '雷鸣地雷': 5}

wb = openpyxl.load_workbook(P, read_only=True, data_only=True, keep_vba=False)
ws = wb['作战记录']
rows = []
for i, row in enumerate(ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True), 2):
    v = list(row) + [None] * 12
    if not v[4]:
        continue
    B, C, D, op, unit = v[1], v[2], v[3], v[4], v[5]
    kind = {'部署': 'deploy', '技能': 'skill', '撤退': 'retreat'}.get(str(op).strip(), str(op))
    rows.append(dict(row=i, B=(int(float(B)) if B not in (None, '') else None),
                     C=int(float(C)) if C not in (None, '') else 0,
                     D=float(D) if D not in (None, '') else None,
                     op=kind, unit=str(unit or '').strip()))

print(f'共 {len(rows)} 条操作，D 非空 {sum(1 for r in rows if r["D"] is not None)} 条，'
      f'B 非空 {sum(1 for r in rows if r["B"] is not None)} 条')
spent = 0.0
print(f'\n{"行":>3} {"操作":<8} {"单位":<10} {"B":>4} {"C":>3} {"D":>6} {"T=D+C/30":>9} '
      f'{"模型读数":>8} {"差":>5}')
mismatch_start = None
for r in rows:
    T = (r['D'] or 0) + r['C'] / FPS
    model = INIT + T - spent
    if r['B'] is not None:
        d = model - r['B']
        if mismatch_start is None and abs(d) > 1.5:
            mismatch_start = r['row']
        print(f'{r["row"]:>3} {r["op"]:<8} {r["unit"]:<10} {r["B"]:>4} {r["C"]:>3} {r["D"]:>6.0f} '
              f'{T:>9.2f} {model:>8.2f} {d:>5.2f}')
    if r['op'] == 'deploy':
        spent += COST.get(r['unit'], 0)
    elif r['op'] == 'retreat':
        spent -= COST.get(r['unit'], 0) * 0.5
print(f'\n首个「模型 vs 读数」偏差 >1.5 的行：{mismatch_start}')
print(f'最终累计消耗 = {spent}')
