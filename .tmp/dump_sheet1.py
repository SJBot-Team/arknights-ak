# -*- coding: utf-8 -*-
"""完整 dump 赛博塑料 h12-4 xlsm 的 Sheet1（疑似作者自己的时间点记录）"""
import sys
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl

P = r'E:\__file\舟数据\关卡\h12-4\赛博塑料h12-4.xlsm'
wb = openpyxl.load_workbook(P, read_only=True, data_only=True, keep_vba=False)
ws = wb['Sheet1']
print(f'Sheet1 维度: {ws.max_row} x {ws.max_column}')
for i, row in enumerate(ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True), 1):
    cells = [('' if x is None else str(x)) for x in row]
    if any(c.strip() for c in cells):
        print(f'{i:>3} | ' + ' | '.join(cells))
