# -*- coding: utf-8 -*-
"""临时：完整 dump 赛博塑料 h12-4 xlsm（所有工作表名 + 作战记录全列 + 设置区）"""
import sys
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl

P = r'E:\__file\舟数据\关卡\h12-4\赛博塑料h12-4.xlsm'
wb = openpyxl.load_workbook(P, read_only=True, data_only=True, keep_vba=False)
print('工作表:', wb.sheetnames)
ws = wb['作战记录']
print('维度:', ws.max_row, 'x', ws.max_column)
for i, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), 1):
    cells = [(f'{chr(65+j)}={str(x)[:18]}' if x is not None else '') for j, x in enumerate(row)]
    print(i, ' | '.join(c for c in cells if c))
print('--- 尾部 ---')
rows = list(ws.iter_rows(min_row=45, max_row=min(ws.max_row, 75), values_only=True))
for i, row in enumerate(rows, 45):
    cells = [(f'{chr(65+j)}={str(x)[:18]}' if x is not None else '') for j, x in enumerate(row)]
    if any(cells):
        print(i, ' | '.join(c for c in cells if c))
for name in wb.sheetnames:
    if name == '作战记录':
        continue
    w2 = wb[name]
    print(f'\n=== {name} ({w2.max_row}x{w2.max_column}) ===')
    for i, row in enumerate(w2.iter_rows(min_row=1, max_row=min(w2.max_row, 25), values_only=True), 1):
        cells = [(str(x)[:22] if x is not None else '') for x in row]
        if any(cells):
            print(i, [c for c in cells if c])
