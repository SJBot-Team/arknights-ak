# -*- coding: utf-8 -*-
"""临时：检查 h12-4 (version 1).xlsb.xlsx 与 快速.xlsm 的结构"""
import sys
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl

for P in (r'E:\__file\舟数据\关卡\h12-4\h12-4 (version 1).xlsb.xlsx',
          r'E:\__file\舟数据\关卡\h12-4\赛博塑料h12-4 - 快速.xlsm'):
    print('=' * 70)
    print(P)
    try:
        wb = openpyxl.load_workbook(P, read_only=True, data_only=True, keep_vba=False)
    except Exception as ex:
        print('  打开失败:', ex)
        continue
    print('  工作表:', wb.sheetnames)
    for name in wb.sheetnames:
        ws = wb[name]
        print(f'  --- {name}: {ws.max_row} x {ws.max_column}')
        shown = 0
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 8), values_only=True), 1):
            cells = [(str(x)[:24] if x is not None else '') for x in row]
            if any(cells):
                print('     ', i, [c for c in cells if c])
                shown += 1
        if name.lower().startswith('作战') or '战斗' in name:
            print('      ... 全表（前 25 行 A..N）')
            for i, row in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 25), values_only=True), 1):
                cells = [(f'{chr(65+j)}={str(x)[:20]}' if x is not None else '') for j, x in enumerate(row[:14])]
                if any(cells):
                    print(f'      {i:>3}', ' | '.join(c for c in cells if c))
