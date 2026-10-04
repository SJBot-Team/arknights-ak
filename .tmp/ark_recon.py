# -*- coding: utf-8 -*-
"""ArknightR reconnaissance - read-only."""
import os, sys, json, re, struct, collections, io
sys.stdout.reconfigure(encoding='utf-8', errors='replace')

R = r'E:\__file\ak\ArknightR'
D = os.path.join(R, 'ArknightR_Data')

def hr(t): print('\n' + '='*20 + ' ' + t + ' ' + '='*20)

# ---------- 1. PlayerSettings from globalgamemanagers ----------
hr('1. UnityPy: globalgamemanagers PlayerSettings')
try:
    import UnityPy
    env = UnityPy.load(os.path.join(D, 'globalgamemanagers'))
    print('unity_version:', env.files[list(env.files)[0]].unity_version if hasattr(env,'files') else '?')
    for o in env.objects:
        if o.type.name == 'PlayerSettings':
            d = o.read_typetree()
            for k in ('productName','companyName','m_ProductName','m_CompanyName','bundleVersion','applicationIdentifier','defaultScreenWidth','defaultScreenHeight','m_ShowUnitySplashScreen','productGUID'):
                if k in d: print(k, '=', repr(d[k]))
            # dump all top-level keys
            print('--- PlayerSettings keys ---')
            print(sorted(d.keys()))
            break
except Exception as e:
    print('ERR', type(e).__name__, e)

# ---------- 2. Addressables catalog ----------
hr('2. Addressables catalog.json')
cat = os.path.join(D, 'StreamingAssets', 'aa', 'catalog.json')
raw = open(cat, 'rb').read()
print('size', len(raw), 'head', raw[:120])
try:
    j = json.loads(raw.decode('utf-8'))
    print('top keys:', list(j.keys())[:40])
    print('m_InternalIds count:', len(j.get('m_InternalIds', [])))
    ii = j.get('m_InternalIds', [])
    print('first 15 internalIds:'); [print('   ', x) for x in ii[:15]]
    print('last 5 internalIds:'); [print('   ', x) for x in ii[-5:]]
    # extension histogram of internal ids
    ext = collections.Counter(os.path.splitext(x)[1] for x in ii)
    print('internalId ext hist:', ext.most_common(15))
    # basename prefixes (for non-bundle ids = asset addresses)
    pre = collections.Counter()
    for x in ii:
        b = os.path.basename(x)
        if not b.endswith('.bundle'):
            pre[re.split(r'[_./]', b)[0]] += 1
    print('non-bundle id prefix hist (top40):', pre.most_common(40))
    loc = j.get('m_locations') or {}
    print('m_locations type', type(loc), 'len', len(loc) if hasattr(loc,'__len__') else '?')
except Exception as e:
    print('catalog parse ERR', type(e).__name__, e)
    print('ascii sample:', raw[:400])

# ---------- 3. Bundle naming analysis ----------
hr('3. Bundle naming / structure')
bundles = []
for root, dirs, files in os.walk(D):
    for f in files:
        if f.endswith('.bundle'):
            bundles.append(os.path.join(root, f))
print('total bundles', len(bundles))
names = [os.path.basename(b) for b in bundles]
base = [re.sub(r'_[0-9a-f]{32}\.bundle$', '', n).replace('.bundle','') for n in names]
c = collections.Counter(base)
print('unique base names', len(c))
print('top 60 by bundle-count:')
for k, v in c.most_common(60):
    print('   %-60s %d' % (k, v))
# group prefixes
pref = collections.Counter(b.split('_')[0] for b in base)
print('prefix hist (top 50):', pref.most_common(50))
# magic of a sample of bundles
mag = collections.Counter()
szs = []
for b in bundles[:1200]:
    with open(b, 'rb') as fh:
        mag[fh.read(8)] += 1
    szs.append(os.path.getsize(b))
print('magic hist (first1200):', mag.most_common(5))
szs.sort()
print('bundle size: min %d  p50 %d  p90 %d  max %d  sum %.1fMB' % (
    szs[0], szs[len(szs)//2], szs[int(len(szs)*0.9)], szs[-1], sum(szs)/1048576))

# ---------- 4. Data files: battle.xlsx / data.sav ----------
hr('4. battle.xlsx')
x = os.path.join(D, 'StreamingAssets', 'battle.xlsx')
print('head bytes', open(x,'rb').read(16).hex())
try:
    import openpyxl
    wb = openpyxl.load_workbook(x, read_only=True, data_only=True)
    print('sheets:', wb.sheetnames)
    for sn in wb.sheetnames[:6]:
        ws = wb[sn]
        print('--- sheet %r dims=%s' % (sn, ws.calculate_dimension()))
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            print('   ', [str(v)[:40] if v is not None else '' for v in row[:14]])
            if i >= 6: break
    wb.close()
except Exception as e:
    print('xlsx ERR', type(e).__name__, e)

hr('5. data.sav')
s = os.path.join(D, 'StreamingAssets', 'data.sav')
raw = open(s, 'rb').read()
print('size', len(raw), 'head64', raw[:64].hex())
print('head repr', repr(raw[:300]))
# ascii strings
asc = re.findall(rb'[\x20-\x7e]{5,}', raw)
print('ascii strings:', [a.decode() for a in asc[:60]])
u16 = re.findall(rb'(?:[\x20-\x7e]\x00){4,}', raw)
print('utf16 strings:', [a.decode('utf-16-le') for a in u16[:40]])

hr('6. settings.json / link.xml / RuntimeInitializeOnLoads.json')
print(open(os.path.join(D,'StreamingAssets','aa','settings.json'), encoding='utf-8').read()[:900])
print('--- link.xml ---')
print(open(os.path.join(D,'StreamingAssets','aa','AddressablesLink','link.xml'), encoding='utf-8').read()[:1500])
print('--- RuntimeInitializeOnLoads.json ---')
print(open(os.path.join(D,'RuntimeInitializeOnLoads.json'), encoding='utf-8').read()[:700])
print('--- ScriptingAssemblies.json ---')
print(open(os.path.join(D,'ScriptingAssemblies.json'), encoding='utf-8').read()[:900])
print('--- app.info ---', repr(open(os.path.join(D,'app.info'),'rb').read()))
print('--- boot.config ---'); print(open(os.path.join(D,'boot.config'), encoding='utf-8', errors='replace').read())
